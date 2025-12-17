import io
from openpyxl import load_workbook


def read_excel_to_dicts(file_bytes):
    """
    Читает Excel-файл и возвращает список словарей.

    Args:
        file_bytes: Байтовое содержимое Excel-файла

    Returns:
        Список словарей вида:
        [
            {'p_name': '...', 'p_key': '...', 'p_guide': '...', 'p_art': '...'},
            ...
        ]
    """
    # Открываем Excel-файл из байтов
    wb = load_workbook(filename=io.BytesIO(file_bytes))
    ws = wb.active  # Берем активный лист

    # Определяем заголовки (первая строка)
    headers = []
    for cell in ws[1]:
        headers.append(str(cell.value).strip() if cell.value else f"column_{cell.column}")

    # Собираем данные
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = {}
        for i, cell_value in enumerate(row):
            if i < len(headers):
                # Преобразуем заголовки Excel в имена полей БД
                header = headers[i]
                if header == 'НАЗВАНИЕ':
                    row_dict['p_name'] = str(cell_value).strip() if cell_value else ''
                elif header == 'КЛЮЧ':
                    row_dict['p_key'] = str(cell_value).strip() if cell_value else ''
                elif header == 'ИНСТРУКЦИЯ':
                    row_dict['p_guide'] = str(cell_value).strip() if cell_value else ''
                elif header == 'АРТИКУЛ':
                    row_dict['p_art'] = str(cell_value).strip() if cell_value else ''

        # Добавляем только если есть данные
        if any(row_dict.values()):
            data.append(row_dict)

    wb.close()
    return data
